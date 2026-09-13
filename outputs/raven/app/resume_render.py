import io
import re
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from docx.shared import Inches, Pt


def _lines(markdown:str):
    return [line.rstrip() for line in markdown.replace("\r","").split("\n")]


def render_pdf(markdown:str)->bytes:
    output=io.BytesIO();styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ResumeName",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=16,leading=18,alignment=TA_CENTER,spaceAfter=5))
    styles.add(ParagraphStyle(name="ResumeHeading",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=10,leading=12,spaceBefore=7,spaceAfter=3,borderWidth=0,borderPadding=0))
    styles.add(ParagraphStyle(name="ResumeBody",parent=styles["BodyText"],fontName="Helvetica",fontSize=9,leading=11,spaceAfter=2))
    styles.add(ParagraphStyle(name="ResumeBullet",parent=styles["BodyText"],fontName="Helvetica",fontSize=9,leading=11,leftIndent=11,firstLineIndent=-7,spaceAfter=1))
    doc=SimpleDocTemplate(output,pagesize=LETTER,rightMargin=.58*inch,leftMargin=.58*inch,topMargin=.48*inch,bottomMargin=.48*inch,title="Targeted Resume",author="RAVEN")
    story=[];first_heading=True
    for raw in _lines(markdown):
        line=raw.strip()
        if not line:story.append(Spacer(1,3));continue
        if line.startswith("#"):
            text=re.sub(r"^#+\s*","",line);style="ResumeName" if first_heading else "ResumeHeading";first_heading=False;story.append(Paragraph(escape(text),styles[style]))
        elif re.match(r"^[-*•]\s+",line):story.append(Paragraph("• "+escape(re.sub(r"^[-*•]\s+","",line)),styles["ResumeBullet"]))
        else:story.append(Paragraph(escape(re.sub(r"\*\*(.+?)\*\*",r"\1",line)),styles["ResumeBody"]))
    doc.build(story);return output.getvalue()


def render_docx(markdown:str)->bytes:
    document=Document();section=document.sections[0];section.top_margin=Inches(.45);section.bottom_margin=Inches(.45);section.left_margin=Inches(.55);section.right_margin=Inches(.55)
    first_heading=True
    for raw in _lines(markdown):
        line=raw.strip()
        if not line:continue
        if line.startswith("#"):
            text=re.sub(r"^#+\s*","",line);p=document.add_paragraph();p.style="Title" if first_heading else "Heading 2";p.add_run(text);first_heading=False
        elif re.match(r"^[-*•]\s+",line):document.add_paragraph(re.sub(r"^[-*•]\s+","",line),style="List Bullet")
        else:document.add_paragraph(re.sub(r"\*\*(.+?)\*\*",r"\1",line))
    for style in document.styles:
        if hasattr(style,"font"):style.font.name="Arial";style.font.size=Pt(10)
    output=io.BytesIO();document.save(output);return output.getvalue()
